from docx import Document as DocxDocument
from langchain_core.documents import Document

def load_docx(file_path):

    doc = DocxDocument(file_path)

    text = "\n".join(
    paragraph.text
    for paragraph in doc.paragraphs
    if paragraph.text.strip()
)

    return [
        Document(
            page_content=text,
            metadata={"source": file_path}
        )
    ]